"""
Resume Parser Module for ByteResume
NoSQL Database Architecture Project
Extracts text from PDF and DOCX files, identifies technical skills,
soft skills, education, work experience, and projects.
"""

import os
import re
from pypdf import PdfReader
import docx

# Comprehensive skill taxonomy for accurate NLP keyword matching
SKILL_TAXONOMY = {
    "languages": [
        "python", "javascript", "typescript", "java", "c++", "c", "c#", "golang", "go",
        "rust", "ruby", "php", "swift", "kotlin", "scala", "r", "dart", "sql", "html", "css"
    ],
    "frameworks": [
        "react", "react.js", "reactjs", "next.js", "nextjs", "vue", "vue.js", "angular",
        "node.js", "nodejs", "express", "express.js", "django", "flask", "fastapi",
        "spring", "spring boot", "asp.net", ".net", "laravel", "rails", "ruby on rails"
    ],
    "databases": [
        "mongodb", "redis", "postgresql", "postgres", "mysql", "sqlite", "cassandra",
        "couchdb", "elasticsearch", "neo4j", "mariadb", "dynamodb", "oracle", "firebase"
    ],
    "tools": [
        "docker", "kubernetes", "git", "github", "gitlab", "aws", "azure", "gcp",
        "google cloud", "linux", "jenkins", "terraform", "ansible", "nginx", "graphql",
        "rest api", "restful", "kafka", "rabbitmq", "postman", "Vite", "webpack"
    ],
    "concepts": [
        "data structures", "algorithms", "machine learning", "deep learning",
        "artificial intelligence", "nlp", "computer vision", "microservices",
        "ci/cd", "system design", "object-oriented programming", "oop",
        "agile", "scrum", "distributed systems", "cloud computing"
    ],
    "soft_skills": [
        "communication", "teamwork", "leadership", "problem solving", "critical thinking",
        "time management", "adaptability", "work ethic", "collaboration", "creatiTechy",
        "attention to detail", "emotional intelligence", "decision making", "presentation",
        "conflict resolution", "mentoring"
    ]
}


def extract_text_from_pdf(file_path):
    """Extract raw text from PDF file using pypdf."""
    text = ""
    try:
        reader = PdfReader(file_path)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    except Exception as e:
        print(f"Error reading PDF {file_path}: {e}")
    return text.strip()


def extract_text_from_docx(file_path):
    """Extract raw text from DOCX file using python-docx."""
    text = ""
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            if para.text:
                text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text += cell.text + " "
                text += "\n"
    except Exception as e:
        print(f"Error reading DOCX {file_path}: {e}")
    return text.strip()


def extract_text(file_path):
    """Detect file type and extract text accordingly."""
    if not os.path.exists(file_path):
        return ""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    else:
        # Fallback to plain text read
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()


def extract_skills(text):
    """Extract technical and soft skills by matching against comprehensive skill taxonomy."""
    text_lower = " " + text.lower() + " "
    found_technical = set()
    found_soft = set()

    for category, skills in SKILL_TAXONOMY.items():
        for skill in skills:
            # Word boundary regex pattern for accurate matching (e.g. avoid 'c' in 'cat')
            escaped = re.escape(skill)
            pattern = rf"(?:\b|(?<=[^a-zA-Z0-9]))" + escaped + rf"(?:\b|(?=[^a-zA-Z0-9]))"
            if re.search(pattern, text_lower):
                # Canonical display name
                canonical = skill.title()
                if skill.lower() in ["sql", "aws", "gcp", "nlp", "oop", "api", "rest api", "ci/cd"]:
                    canonical = skill.upper()
                elif skill.lower() in ["mongodb", "redis", "postgresql", "mysql", "sqlite", "graphql"]:
                    canonical = skill.capitalize()
                    if skill.lower() == "mongodb": canonical = "MongoDB"
                    if skill.lower() == "postgresql": canonical = "PostgreSQL"
                    if skill.lower() == "mysql": canonical = "MySQL"
                    if skill.lower() == "graphql": canonical = "GraphQL"
                elif skill.lower() in ["react", "react.js", "reactjs"]:
                    canonical = "React"
                elif skill.lower() in ["node.js", "nodejs"]:
                    canonical = "Node.js"
                elif skill.lower() in ["next.js", "nextjs"]:
                    canonical = "Next.js"
                elif skill.lower() in ["vue", "vue.js"]:
                    canonical = "Vue.js"

                if category == "soft_skills":
                    found_soft.add(canonical)
                else:
                    found_technical.add(canonical)

    return sorted(list(found_technical)), sorted(list(found_soft))


def extract_sections(text):
    """Extract education, experience, and project entries from resume text using heuristics."""
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    education = []
    experience = []
    projects = []

    current_section = None
    edu_keywords = ["education", "academic", "qualification", "degree", "university", "college", "school"]
    exp_keywords = ["experience", "employment", "work history", "internship", "work experience", "career"]
    proj_keywords = ["projects", "personal projects", "academic projects", "key projects"]

    for line in lines:
        line_lower = line.lower()
        # Check if line is a section header (short line matching header keyword)
        if len(line.split()) <= 4:
            if any(k in line_lower for k in edu_keywords):
                current_section = "education"
                continue
            elif any(k in line_lower for k in exp_keywords):
                current_section = "experience"
                continue
            elif any(k in line_lower for k in proj_keywords):
                current_section = "projects"
                continue

        if current_section == "education":
            if len(education) < 5 and len(line) > 3:
                education.append(line)
        elif current_section == "experience":
            if len(experience) < 8 and len(line) > 3:
                experience.append(line)
        elif current_section == "projects":
            if len(projects) < 6 and len(line) > 3:
                projects.append(line)

    # If heuristic section extraction was sparse, run regex pattern fallbacks
    if not education:
        edu_patterns = re.findall(r"(?:B\.?Tech|B\.?E\.?|M\.?Tech|M\.?S\.?|B\.?Sc|M\.?Sc|Bachelor|Master|Ph\.?D)[^\n.,;]*", text, re.IGNORECASE)
        education = [e.strip() for e in edu_patterns[:4] if len(e.strip()) > 3]

    if not experience:
        exp_patterns = re.findall(r"(?:Engineer|Developer|Intern|Consultant|Analyst|Lead|Specialist)[^\n.,;]*", text, re.IGNORECASE)
        experience = [e.strip() for e in exp_patterns[:5] if len(e.strip()) > 3]

    if not projects:
        proj_patterns = re.findall(r"(?:Project|System|Application|Platform|Dashboard|Analyzer|App):?\s*([^\n]+)", text, re.IGNORECASE)
        projects = [p.strip() for p in proj_patterns[:4] if len(p.strip()) > 3]

    return education, experience, projects


def parse_resume(file_path=None, raw_text=None, file_name=None):
    """Main function to parse resume file or raw text."""
    if file_path and os.path.exists(file_path):
        text = extract_text(file_path)
        name = file_name or os.path.basename(file_path)
        ext = os.path.splitext(name)[1].lower().replace(".", "")
    elif raw_text:
        text = raw_text
        name = file_name or "uploaded_resume.txt"
        ext = "txt"
    else:
        text = ""
        name = file_name or "unknown"
        ext = "unknown"

    tech_skills, soft_skills = extract_skills(text)
    education, experience, projects = extract_sections(text)

    all_skills = sorted(list(set(tech_skills + soft_skills)))

    return {
        "fileName": name,
        "fileType": ext,
        "rawText": text,
        "skills": all_skills,
        "technicalSkills": tech_skills,
        "softSkills": soft_skills,
        "education": education,
        "experience": experience,
        "projects": projects
    }
